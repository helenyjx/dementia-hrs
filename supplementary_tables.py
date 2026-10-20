"""Formatted manuscript tables (Tables 2 and 3) and Supplementary Tables S10-S19, from results/tables/ and the
outcome-status table. Writes one CSV per table to results/tables/supplementary/ and, with --docx, a Word file.

Table 2   internal validation (all participants)
Table 3   temporal validation (all participants)
S10 outcome status by cohort and horizon
S11 participants with complete vs incomplete predictors
S12 cohort overlap and temporal-validation performance by overlap group
S13 accuracy (all participants) and complete-case results
S14 paired differences between models, internal validation
S15 paired differences between models, temporal validation
S16 selected hyperparameters and decision thresholds
S17 runtime
S18 SHAP vs permutation importance
S19 outcome-definition sensitivity analysis

Usage: python supplementary_tables.py [--docx "paper/Supplementary Tables.docx"]
"""
import argparse
import json
import os
from collections import Counter

import numpy as np
import pandas as pd

from make_figures import MODELS, NAMES, VARIABLE_LABELS

ROOT = os.path.dirname(os.path.abspath(__file__))
TABLES = os.path.join(ROOT, 'results', 'tables')
OUT = os.path.join(TABLES, 'supplementary')
STATUS = os.path.join(ROOT, 'data', 'preprocessed data', 'outcome_status.csv')
COHORT = {'temporal': '2010 cohort', 'internal': '2000 cohort'}
PCT = ('AUROC', 'AUPRC', 'Accuracy', 'Precision', 'Recall', 'F1')
MAIN_METRICS = ['AUROC', 'AUPRC', 'Precision', 'Recall', 'F1', 'Brier', 'calibration_intercept',
                'calibration_slope', 'OE_ratio']
METRIC_LABEL = {'AUROC': 'AUROC (%)', 'AUPRC': 'AUPRC (%)', 'Accuracy': 'Accuracy (%)', 'Precision': 'Precision (%)',
                'Recall': 'Recall (%)', 'F1': 'F1 (%)', 'Brier': 'Brier score', 'calibration_intercept':
                'Calibration intercept', 'calibration_slope': 'Calibration slope', 'OE_ratio': 'O/E ratio'}


def read(name):
    return pd.read_csv(os.path.join(TABLES, name))


def fmt(est, lo=None, hi=None, metric='AUROC'):
    if metric in PCT:
        s, f = 100, '{:.1f}'
    elif metric == 'Brier':
        s, f = 1, '{:.3f}'
    else:
        s, f = 1, '{:.2f}'
    out = f.format(est * s)
    if lo is not None:
        out += f' ({f.format(lo * s)} to {f.format(hi * s)})'
    return out


def ordered(df, col='model'):
    return df.assign(_o=df[col].map({m: i for i, m in enumerate(MODELS)})).sort_values(
        [c for c in ('horizon', 'h') if c in df] + ['_o']).drop(columns='_o')


def metric_table(df, metrics, population='all'):
    """Wide table: one row per horizon and model, one column per metric, 'estimate (95% CI)'."""
    df = ordered(df[df['population'] == population])
    rows = []
    for (h, m), g in df.groupby(['horizon', 'model'], sort=False):
        g = g.set_index('metric')
        n = g['n'].iloc[0] if 'n' in g else g['n_mean'].iloc[0]
        cases = g['cases'].iloc[0] if 'cases' in g else g['cases_mean'].iloc[0]
        row = {'Horizon': f'{h}-year', 'Model': NAMES[m], 'n': f'{n:,.0f}', 'Cases': f'{cases:,.0f}',
               'Prevalence (%)': f'{100 * cases / n:.1f}'}
        for k in metrics:
            row[METRIC_LABEL[k]] = fmt(g.at[k, 'estimate'], g.at[k, 'ci_low'], g.at[k, 'ci_high'], k)
        rows.append(row)
    return pd.DataFrame(rows)


def t10_outcome_status():
    s = pd.read_csv(STATUS)
    label = {'demcls00-06': ('2000', '6-year (2000-2006)'), 'demcls00-08': ('2000', '8-year (2000-2008)'),
             'demcls10-16': ('2010', '6-year (2010-2016)'), 'demcls10-18': ('2010', '8-year (2010-2018)')}
    rows = []
    for _, r in s.iterrows():
        cohort, window = label[r['window']]
        e = r['eligible']
        pct = lambda v: f'{v:,} ({100 * v / e:.1f})'
        rows.append({'Baseline cohort': cohort, 'Horizon (window)': window, 'Eligible at baseline': f'{e:,}',
                     'Incident dementia, n (%)': pct(r['incident dementia']),
                     'Assessed dementia-free at final wave, n (%)': pct(r['assessed dementia-free']),
                     'Died before final wave, no dementia observed, n (%)':
                         pct(r['died before final wave, no dementia observed']),
                     'Alive or unknown, not assessed at final wave, n (%)':
                         pct(r['alive or unknown, not assessed at final wave']),
                     'Analytic cohort': f'{r["incident dementia"] + r["assessed dementia-free"]:,}'})
    return pd.DataFrame(rows)


def t11_complete_vs_incomplete():
    d = read('complete_vs_incomplete.csv')
    return pd.DataFrame({
        'Cohort': d['cohort'], 'Horizon': d['horizon (years)'].map(lambda h: f'{h}-year'),
        'Characteristic': d['characteristic'],
        'Complete predictors': [f'{a:.2f}' if max(a, b) < 10 else f'{a:.1f}' for a, b in zip(d['complete predictors'], d['incomplete predictors'])],
        'Incomplete predictors': [f'{b:.2f}' if max(a, b) < 10 else f'{b:.1f}' for a, b in zip(d['complete predictors'], d['incomplete predictors'])],
        'Standardized difference': d['standardized difference'].map(lambda v: f'{v:.2f}'),
        'n complete / incomplete': [f'{a:,} / {b:,}' for a, b in zip(d['n complete'], d['n incomplete'])]})


def t12_overlap():
    o = read('cohort_overlap.csv')
    b = ordered(read('temporal_by_overlap.csv'))
    groups = ['returning participant', 'new participant, development household', 'new participant and household']
    rows = []
    for _, r in o.iterrows():
        h, n = r['horizon (years)'], r['2010 analytic cohort']
        for g in groups:
            first = b[(b['h'] == h) & (b['group'] == g)].iloc[0]
            row = {'Horizon': f'{h}-year', 'Group': g[0].upper() + g[1:],
                   'n (% of cohort)': f'{r[g]:,} ({100 * r[g] / n:.1f})',
                   'Cases (%)': f'{first["cases"]:,} ({100 * first["cases"] / first["n"]:.1f})'}
            for _, m in b[(b['h'] == h) & (b['group'] == g)].iterrows():
                row[f'{NAMES[m["model"]]} AUROC / AUPRC (%)'] = f'{m["AUROC"]:.1f} / {m["AUPRC"]:.1f}'
            rows.append(row)
        rows.append({'Horizon': f'{h}-year', 'Group': 'Households in 2010 cohort, of which also in 2000 cohort',
                     'n (% of cohort)': f'{r["households in 2010 cohort"]:,}, of which {r["of which also in 2000 cohort"]:,} '
                                        f'({100 * r["of which also in 2000 cohort"] / r["households in 2010 cohort"]:.1f})'})
    return pd.DataFrame(rows).fillna('')


def t13_accuracy_complete():
    i, t = read('internal_metrics.csv'), read('temporal_metrics.csv')
    metrics = ['AUROC', 'AUPRC', 'Accuracy', 'Precision', 'Recall', 'F1', 'Brier', 'calibration_slope', 'OE_ratio']
    acc, complete = [], []
    for setting, d in (('Internal (2000 cohort)', i), ('Temporal (2010 cohort)', t)):
        a = metric_table(d, ['Accuracy'])
        a.insert(0, 'Validation', setting)
        acc.append(a)
        c = metric_table(d, metrics, population='complete')
        c.insert(0, 'Validation', setting)
        complete.append(c)
    return pd.concat(acc, ignore_index=True), pd.concat(complete, ignore_index=True)


def pairwise_table(name):
    d = read(name)
    d = d[d['population'] == 'all']
    rows = []
    for _, r in d.iterrows():
        a, b = r['model_a'], r['model_b']
        # orient each pair so that model A comes first in the fixed model order (higher 6-year internal AUROC)
        base_b = r['difference'] / (r['relative_difference_pct'] / 100)  # comparator (model_b) value
        diff, lo, hi = r['difference'], r['ci_low'], r['ci_high']
        if MODELS.index(a) > MODELS.index(b):
            a, b = b, a
            base_b = base_b + diff  # the new comparator is the former model_a
            diff, lo, hi = -diff, -hi, -lo
        p = r['p_two_sided']
        rows.append({'horizon': r['horizon'], 'oa': MODELS.index(a), 'ob': MODELS.index(b), 'metric': r['metric'],
                     'Horizon': f'{r["horizon"]}-year', 'Model A': NAMES[a], 'Model B': NAMES[b],
                     'Metric': METRIC_LABEL[r['metric']].replace(' (%)', ''),
                     'Absolute difference, A - B (95% CI)': (
                         f'{100 * diff:.1f} ({100 * lo:.1f} to {100 * hi:.1f}) points' if r['metric'] != 'Brier'
                         else f'{diff:.4f} ({lo:.4f} to {hi:.4f})'),
                     'Relative difference (%)': f'{100 * diff / base_b:.1f}',
                     'p': '<0.001' if p < 0.001 else f'{p:.3f}'})
    order = {'AUROC': 0, 'AUPRC': 1, 'F1': 2, 'Brier': 3}
    out = pd.DataFrame(rows)
    out = out.assign(om=out['metric'].map(order)).sort_values(['horizon', 'om', 'oa', 'ob'])
    return out.drop(columns=['horizon', 'oa', 'ob', 'metric', 'om'])


def _params(p):
    return ', '.join(f'{k}={v}' for k, v in p.items()) if p else '(defaults)'


def t16_selection():
    rows = []
    for h in (6, 8):
        for m in MODELS:
            s = json.load(open(os.path.join(ROOT, 'results', f'h{h}', m, 'selection.json')))
            reps = s['per_repeat']
            counts = Counter(_params(r['params']) for r in reps)
            thr = np.array([r['threshold'] for r in reps])
            rows.append({'Horizon': f'{h}-year', 'Model': NAMES[m],
                         'Configurations selected in internal repeats (times of 10)':
                             '; '.join(f'{c} [{n}]' for c, n in counts.most_common()),
                         'Internal thresholds, median (range)': f'{np.median(thr):.3f} ({thr.min():.3f} to {thr.max():.3f})',
                         'Temporal-validation configuration': _params(s['temporal']['params']),
                         'Temporal-validation threshold': f'{s["temporal"]["threshold"]:.3f}'})
    return pd.DataFrame(rows)


def t17_runtime():
    d = ordered(read('runtime.csv'))
    return pd.DataFrame({'Horizon': d['horizon'].map(lambda h: f'{h}-year'), 'Model': d['model'].map(NAMES),
                         'Fits': d['fits'], 'Median seconds per fit': d['median_seconds_per_fit'].map(lambda v: f'{v:.2f}'),
                         'Total minutes': d['total_minutes'].map(lambda v: f'{v:.1f}')})


def t18_shap():
    perm, shap = read('importance_permutation_h6.csv'), read('importance_shap_h6.csv')
    agree = ordered(read('importance_rank_agreement_h6.csv'))
    a = pd.DataFrame({'Model': agree['model'].map(NAMES),
                      'Spearman rank correlation (60 predictors)': agree['spearman_rank_correlation'].map(lambda v: f'{v:.2f}'),
                      'Shared top-10 predictors': agree['shared_top10']})
    rows = []
    for m in agree['model']:
        p = perm[perm['model'] == m].set_index('variable')
        s = shap[shap['model'] == m].set_index('variable')
        top = sorted(set(p.index[p['rank'] <= 10]) | set(s.index[s['rank'] <= 10]), key=lambda v: (p.at[v, 'rank'], s.at[v, 'rank']))
        for v in top:
            rows.append({'Model': NAMES[m], 'Predictor': VARIABLE_LABELS.get(v, v),
                         'Permutation rank': p.at[v, 'rank'],
                         'Permutation importance, AUPRC drop (95% CI), points': fmt(p.at[v, 'importance_AUPRC_drop'],
                                                                                   p.at[v, 'ci_low'], p.at[v, 'ci_high'], 'AUPRC'),
                         'SHAP rank': s.at[v, 'rank'], 'Mean |SHAP| (log-odds)': f'{s.at[v, "mean_abs_SHAP"]:.3f}'})
    return a, pd.DataFrame(rows)


def t19_sensitivity():
    out = []
    metrics = ['AUROC', 'AUPRC', 'F1', 'Brier']
    for setting, prim, sens in (('Internal (2000 cohort)', 'internal_metrics.csv', 'sensitivity_any_internal_metrics.csv'),
                                ('Temporal (2010 cohort)', 'temporal_metrics.csv', 'sensitivity_any_temporal_metrics.csv')):
        s = read(sens)
        p = read(prim)
        p = p[p['model'].isin(s['model'].unique())]
        for label, d in (('Assessed dementia-free at any follow-up wave (sensitivity)', s),
                         ('Assessed dementia-free at the final wave (primary)', p)):
            w = metric_table(d, metrics)
            w.insert(0, 'Non-case definition', label)
            w.insert(0, 'Validation', setting)
            out.append(w)
    out = pd.concat(out, ignore_index=True)
    return out.sort_values(['Validation', 'Horizon', 'Model'], key=lambda c: c.map(
        {NAMES[m]: i for i, m in enumerate(MODELS)}) if c.name == 'Model' else c, kind='stable')


TITLES = {
    'Table2': 'Table 2. Internal validation in the 2000 development cohort (all participants). Mean over 10 random '
              '60/20/20 splits with 95% CIs from the corrected resampled t interval; n and cases are per test set.',
    'Table3': 'Table 3. Temporal validation within the HRS (2010 cohort, all participants). Models fitted on the full '
              '2000 cohort; 95% CIs from a household cluster bootstrap (2,000 resamples).',
    'S10': 'Supplementary Table S10. Outcome status of participants eligible at baseline, by cohort and horizon.',
    'S11': 'Supplementary Table S11. Characteristics of analytic-cohort participants with complete and incomplete '
           'predictors. Values are means or percentages; standardized difference = (incomplete - complete) / pooled SD.',
    'S12': 'Supplementary Table S12. Overlap between the 2010 temporal-validation cohort and the 2000 development '
           'cohort, and temporal-validation performance by overlap group.',
    'S13a': 'Supplementary Table S13a. Accuracy (all participants) at the F1-maximizing threshold selected on '
            'validation data (internal) or carried over from development (temporal).',
    'S13b': 'Supplementary Table S13b. Performance among participants with complete predictors (complete-case '
            'sensitivity analysis).',
    'S14': 'Supplementary Table S14. Paired differences between models, internal validation (all participants). '
           'Corrected resampled t intervals and p values; relative difference with respect to model B.',
    'S15': 'Supplementary Table S15. Paired differences between models, temporal validation (all participants). '
           'Household cluster-bootstrap intervals and p values; relative difference with respect to model B.',
    'S16': 'Supplementary Table S16. Selected hyperparameters and decision thresholds (selection by validation AUPRC; '
           'threshold maximizing F1 on the validation set; temporal threshold = median of the internal thresholds).',
    'S17': 'Supplementary Table S17. Runtime: all fits of the hyperparameter search and the final refits (median '
           'seconds per fit includes prediction on the validation set; neural models on an NVIDIA GeForce RTX 4080).',
    'S18a': 'Supplementary Table S18a. Agreement between SHAP and permutation importance ranks (2010 cohort, 6-year horizon).',
    'S18b': 'Supplementary Table S18b. Predictors in the top 10 by either method.',
    'S19': 'Supplementary Table S19. Outcome-definition sensitivity analysis: non-cases defined as assessed '
           'dementia-free at any follow-up wave, compared with the primary definition.',
}


def build():
    i, t = read('internal_metrics.csv'), read('temporal_metrics.csv')
    s13a, s13b = t13_accuracy_complete()
    s18a, s18b = t18_shap()
    return {'Table2': metric_table(i, MAIN_METRICS), 'Table3': metric_table(t, MAIN_METRICS),
            'S10': t10_outcome_status(), 'S11': t11_complete_vs_incomplete(), 'S12': t12_overlap(),
            'S13a': s13a, 'S13b': s13b, 'S14': pairwise_table('internal_pairwise.csv'),
            'S15': pairwise_table('temporal_pairwise.csv'), 'S16': t16_selection(), 'S17': t17_runtime(),
            'S18a': s18a, 'S18b': s18b, 'S19': t19_sensitivity()}


def write_docx(tables, path):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Cm(1.5))
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(8)
    for k, df in tables.items():
        doc.add_paragraph().add_run(TITLES[k]).bold = True
        tab = doc.add_table(rows=1, cols=len(df.columns))
        tab.style = 'Table Grid'
        for c, col in enumerate(df.columns):
            run = tab.rows[0].cells[c].paragraphs[0].add_run(str(col))
            run.bold = True
        for _, r in df.iterrows():
            cells = tab.add_row().cells
            for c, v in enumerate(r):
                cells[c].text = str(v)
        for row in tab.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs:
                        run.font.size = Pt(7)
        doc.add_paragraph()
    doc.save(path)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--docx', help='also write all tables to this Word file')
    args = ap.parse_args()
    tables = build()
    os.makedirs(OUT, exist_ok=True)
    for k, df in tables.items():
        df.to_csv(os.path.join(OUT, f'{k}.csv'), index=False)
        print(f'{k}: {len(df)} rows')
    if args.docx:
        write_docx(tables, args.docx)
        print('wrote', args.docx)
